import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import Form, FormResponse
from app.services.responses import display_value
from app.services.uploads import original_name


def build_table(db: Session, form: Form) -> tuple[list[str], list[list[str]]]:
    headers = ["Response ID", "Submitted At", "Submitted By"] + [f.label for f in form.fields]
    responses = db.scalars(
        select(FormResponse)
        .where(FormResponse.form_id == form.id)
        .options(selectinload(FormResponse.details), selectinload(FormResponse.user))
        .order_by(FormResponse.id)
    ).all()
    rows = []
    for r in responses:
        values = {d.field_id: json.loads(d.value) for d in r.details}
        row = [
            str(r.id),
            r.created_at.strftime("%Y-%m-%d %H:%M"),
            r.user.full_name if r.user else "Anonymous",
        ]
        for f in form.fields:
            v = values.get(f.id)
            row.append(original_name(v) if f.field_type == "file" and v else display_value(v))
        rows.append(row)
    return headers, rows


def to_excel(headers: list[str], rows: list[list[str]], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1976D2")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows]) if rows else len(h)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 12), 50)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(headers: list[str], rows: list[list[str]], title: str) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small")
    small.fontSize = 7
    small.leading = 9
    head = small.clone("head")
    head.textColor = colors.white
    head.fontName = "Helvetica-Bold"

    def cell(text: str, style) -> Paragraph:
        safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(safe, style)

    data = [[cell(h, head) for h in headers]] + [[cell(c, small) for c in r] for r in rows]
    if not rows:
        data.append([cell("No responses yet", small)] + [cell("", small)] * (len(headers) - 1))
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1976D2")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FB")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    doc.build([Paragraph(f"{title} - Responses", styles["Title"]), Spacer(1, 8), table])
    return buf.getvalue()


def render(db: Session, form: Form, fmt: str) -> tuple[bytes, str, str]:
    """Returns (content, media_type, filename)."""
    headers, rows = build_table(db, form)
    safe = "".join(c if c.isalnum() else "_" for c in form.title)[:40]
    if fmt == "xlsx":
        return (
            to_excel(headers, rows, form.title),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{safe}_responses.xlsx",
        )
    return to_pdf(headers, rows, form.title), "application/pdf", f"{safe}_responses.pdf"
